"""Deterministic export of extracted JSON records into the supplied Excel layout."""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date
import json
import os
from pathlib import Path
import re
import tempfile
import threading
import unicodedata
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "excel_output_template.xlsx"
DEFAULT_EXPORT_PATH = PROJECT_ROOT / "outputs" / "base_llm" / "formatted_data.xlsx"
SHEET_NAME = "データ入力表_temp_202405"
HEADER_ROW = 2
LABEL_ROW = 3
DATA_START_ROW = 4
UNMAPPED_SHEET_NAME = "未対応データ"
_EXPORT_LOCK = threading.RLock()


# Stable IDs are preferred. These aliases cover labels and IDs currently emitted
# by the questionnaire extraction prompts; the workbook itself remains the
# source of truth for column order and displayed field names.
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("測定日", "記入日", "measurement date"),
    "date of birth": ("生年月日", "生年月", "生まれた年", "date of birth"),
    "family": ("お住まいの家族構成", "家族構成"),
    "area": ("お住まいの地域",),
    "momentum": ("コロナ前後の運動量", "新型コロナ前後の運動量"),
    "subjective health": ("普段自分の健康状態をどのように感じていますか", "主観的健康感"),
    "EH1": ("定期的にウォーキングや散歩をしていますか",),
    "EH1-2": ("ウォーキング",),
    "EH1-3": ("散歩", "活動選択散歩"),
    "EH1times": ("ウォーキングや散歩の頻度",),
    "EH1hours": ("ウォーキングや散歩の時間",),
    "EH2": ("定期的にランニングやサイクリングをしていますか",),
    "EH2-2": ("ランニング",),
    "EH2-3": ("サイクリング",),
    "EH2times": ("ランニングやサイクリングの頻度",),
    "EH2hours": ("ランニングやサイクリングの時間",),
    "EH3": ("筋力トレーニング",),
    "EH3times": ("筋力トレーニングの頻度",),
    "EH3hours": ("筋力トレーニングの時間",),
    "EH4": ("その他の運動", "体を動かす習慣"),
    "lococlass1": ("ロコモ予防の運動教室に参加された経験はありますか",),
    "lococlass2": ("現在の方法",),
    "lococlass3": ("現在の頻度",),
    "lococlass4": ("実施頻度の増減",),
    "lococlass5-1": ("運動内容がきつい",),
    "lococlass5-2": ("他の運動を行っているからよい",),
    "lococlass5-3": ("予防運動の必要性をあまり感じない",),
    "lococlass5-4": ("運動に飽きてしまった",),
    "lococlass5-5": ("持病やけがのため",),
    "lococlass5-6": ("多忙のため",),
    "lococlass5-7": ("一人で行うのは難しい",),
    "lococlass5-8": ("ロコモ予防の運動教室のその他の理由",),
    "smoking": ("喫煙習慣はありますか",),
    "drinking": ("飲酒習慣はありますか",),
    "sleeping time": ("平均睡眠時間",),
    "Number of meals": ("食事回数", "食事の回数"),
    "morning meal": ("朝食", "朝食はとりますか"),
    "outings": ("1週間に何回外出しますか", "1週間に何回外出するか"),
    "conversations large": ("会話する人数多い日", "会話する人数平均多い", "1日に何人と会話しますか多い日"),
    "conversations few": ("会話する人数少ない日", "会話する人数少ない", "1日に何人と会話しますか少ない日"),
    "occupation": ("20歳から50歳で最も長く務めた職業",),
    "5Ditch1": ("かゆみを感じる時間はありましたか",),
    "5Ditch2": ("1日に何時間かゆみがありましたか",),
    "5Ditch3": ("かゆみの強さ", "かゆみの強さの点数"),
    "5Ditch4(1)": ("最近の2週間にあなたのかゆみが以下の活動に与える影響について該当する項目の下の空欄に○をつけてください睡眠",),
    "5Ditch4(2)": ("最近の2週間にあなたのかゆみが以下の活動に与える影響について該当する項目の下の空欄に○をつけてください余暇社会的活動",),
    "5Ditch4(3)": ("最近の2週間にあなたのかゆみが以下の活動に与える影響について該当する項目の下の空欄に○をつけてください家事雑用",),
    "5Ditch4(4)": ("最近の2週間にあなたのかゆみが以下の活動に与える影響について該当する項目の下の空欄に○をつけてください仕事家業",),
    "SARCF1": ("4.5㎏の荷物を持ち運び", "4.5kgの荷物を持ち運び"),
    "SARCF2": ("部屋の中を移動",),
    "SARCF3": ("椅子やベッドから立ち上がる",),
    "SARCF4": ("階段を10段のぼる",),
    "SARCF5": ("過去1年間で何回転倒しましたか",),
}

QUESTION_ID_ALIASES: dict[str, set[str]] = {
    "EH1": {"page_004_q1_main", "page_004_field_003"},
    "EH1-2": {"page_004_q1_activity_1", "page_004_table_001_row_001"},
    "EH1-3": {"page_004_q1_activity_2", "page_004_table_001_row_002"},
    "EH1times": {
        "page_004_q1_frequency", "page_004_field_006", "page_004_table_002_row_001",
    },
    "EH1hours": {
        "page_004_q1_duration", "page_004_field_007", "page_004_table_002_row_002",
    },
    "EH2": {"page_004_q2_main", "page_004_field_008"},
    "EH2-2": {"page_004_q2_activity_1", "page_004_table_003_row_001"},
    "EH2-3": {"page_004_q2_activity_2", "page_004_table_003_row_002"},
    "EH2times": {
        "page_004_q2_frequency", "page_004_field_011", "page_004_table_004_row_001",
    },
    "EH2hours": {
        "page_004_q2_duration", "page_004_field_012", "page_004_table_004_row_002",
    },
    "EH3": {"page_004_q3_main", "page_004_field_013"},
    "EH3times": {
        "page_004_q3_frequency", "page_004_field_014", "page_004_table_005_row_001",
    },
    "EH3hours": {
        "page_004_q3_duration", "page_004_field_015", "page_004_table_005_row_002",
    },
    "EH4": {"page_004_q4_text", "page_004_field_016"},
}


def _normalise(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    text = re.sub(r"^(?:問|質問)\s*\d+\s*", "", text)
    return re.sub(r"[^\w\u3040-\u30ff\u3400-\u9fff]+", "", text)


def _natural_key(value: str) -> tuple[tuple[int, Any], ...]:
    return tuple((0, int(part)) if part.isdigit() else (1, part.casefold())
                 for part in re.split(r"(\d+)", value))


def _cell_text(value: Any) -> Any:
    if value is None or value == "":
        return None
    if isinstance(value, date):
        return value
    if isinstance(value, (str, int, float, bool)):
        text = value
    elif isinstance(value, list):
        text = "\n".join(str(item) for item in value if item is not None and str(item).strip())
    elif isinstance(value, dict):
        text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    else:
        text = str(value)
    if isinstance(text, str) and len(text) > 32767:
        raise ValueError("回答がExcelのセル文字数上限を超えています。Excel出力を停止しました。")
    return text


def _answer_status(answer: dict[str, Any]) -> str:
    return str(answer.get("status") or answer.get("answer_status") or "").strip().lower()


def _answer_value(answer: dict[str, Any]) -> Any:
    status = _answer_status(answer)
    if status in {"not_answered", "skipped"}:
        return None
    return answer.get("value")


def _excel_date(value: Any) -> Any:
    text = str(value or "").strip()
    match = re.fullmatch(r"(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})日?", text)
    if not match:
        match = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", text)
    if not match:
        return value
    try:
        return date(*(int(part) for part in match.groups()))
    except ValueError:
        return value


def _date_component(label: str, *, birth: bool) -> str | None:
    normalized_label = _normalise(label)
    if birth and "生年月" in normalized_label:
        return None
    prefixes = ("生まれた",) if birth else ("記入日", "測定日")
    component_by_suffix = {"年": "year", "月": "month", "日": "day"}
    for prefix in prefixes:
        normalized_prefix = _normalise(prefix)
        if normalized_prefix not in normalized_label:
            continue
        suffix = normalized_label.split(normalized_prefix, 1)[1]
        return component_by_suffix.get(suffix)
    return None


def _record_id(record: dict[str, Any], source_path: Path) -> str:
    return str(record.get("person_id") or source_path.stem)


def _is_context_answer(answer: dict[str, Any]) -> bool:
    if answer.get("field_role") == "document_context":
        return True
    question_id = str(answer.get("question_id") or "")
    if re.search(r"_(?:title|footer|(?:sub)?section_header)(?:_\d+)?$", question_id):
        return True
    label = _normalise(answer.get("label_japanese"))
    return label in {
        _normalise(item) for item in (
            "タイトル", "説明文", "ページ番号", "フッター", "アンケート協力依頼",
            "ご協力ありがとうございました", "かゆみスケール",
        )
    } and _answer_value(answer) in (None, "", answer.get("label_japanese"))


def _answer_direct_match(answer: dict[str, Any], field_key: str) -> bool:
    wanted = _normalise(field_key)
    return any(
        _normalise(answer.get(candidate)) == wanted
        for candidate in ("normalized_key", "question_id")
        if answer.get(candidate) not in (None, "")
    )


def _answer_qid_match(answer: dict[str, Any], field_key: str) -> bool:
    question_id = str(answer.get("question_id") or "")
    return question_id in QUESTION_ID_ALIASES.get(field_key, set())


def _field_label_match(answer: dict[str, Any], field_key: str, field_label: str) -> bool:
    source = _normalise(answer.get("label_japanese"))
    if not source:
        return False
    if field_key == "lococlass5-8" and source == _normalise("その他"):
        return False
    aliases = (field_label, *FIELD_ALIASES.get(field_key, ()))
    normalized_aliases = {_normalise(alias) for alias in aliases if _normalise(alias)}
    if source in normalized_aliases:
        return True
    # Permit a longer printed question to wrap a concise template label. Short
    # aliases are kept exact to avoid matching generic words such as "時間".
    return any(len(alias) >= 5 and (alias in source or source in alias)
               for alias in normalized_aliases)


def _find_field(answer: dict[str, Any], fields: list[tuple[str, str]]) -> str | None:
    direct = [key for key, _ in fields if _answer_direct_match(answer, key)]
    if len(direct) == 1:
        return direct[0]
    if len(direct) > 1:
        return None
    by_id = [key for key, _ in fields if _answer_qid_match(answer, key)]
    if len(by_id) == 1:
        return by_id[0]
    if len(by_id) > 1:
        return None
    matches = [key for key, label in fields if _field_label_match(answer, key, label)]
    return matches[0] if len(matches) == 1 else None


def _selected_region_values(answer: dict[str, Any], fields: list[tuple[str, str]]) -> dict[str, Any] | None:
    qid = str(answer.get("question_id") or "")
    label = _normalise(answer.get("label_japanese"))
    is_region_question = qid == "page_011_field_010" or "かゆみがあった部位" in label
    if not is_region_question:
        return None
    raw = answer.get("value")
    selected = raw if isinstance(raw, list) else [raw]
    selected_text = [(_normalise(item), _cell_text(item)) for item in selected if item not in (None, "")]
    result: dict[str, Any] = {}
    for key, field_label in fields:
        if not key.startswith("5Ditch5("):
            continue
        target = _normalise(field_label)
        aliases = [target, *(_normalise(item) for item in FIELD_ALIASES.get(key, ()))]
        for option, original in selected_text:
            if any(alias and (alias == option or alias in option or option in alias)
                   for alias in aliases):
                result[key] = original
                break
    return result


def _combine_component_date(
    answers: list[dict[str, Any]], *, birth: bool,
) -> tuple[str | None, set[int]]:
    tokens = ("生まれた", "生年月") if birth else ("記入日", "測定日")
    parts: dict[str, str] = {}
    consumed: set[int] = set()
    for index, answer in enumerate(answers):
        label = str(answer.get("label_japanese") or "")
        normalized_label = _normalise(label)
        if not any(_normalise(token) in normalized_label for token in tokens):
            continue
        if birth and ("生年月" in normalized_label or "生年月日" in normalized_label):
            value = _cell_text(_answer_value(answer))
            if value not in (None, ""):
                consumed.add(index)
                return _excel_date(value), consumed
            continue
        if not birth and normalized_label in {_normalise("測定日"), _normalise("記入日")}:
            value = _cell_text(_answer_value(answer))
            if value not in (None, ""):
                consumed.add(index)
                return _excel_date(value), consumed
            continue
        component = _date_component(label, birth=birth)
        if not component:
            value = _cell_text(_answer_value(answer))
            if value not in (None, ""):
                consumed.add(index)
                return _excel_date(value), consumed
            continue
        value = _cell_text(_answer_value(answer))
        if value in (None, ""):
            continue
        if component and component not in parts:
            parts[component] = re.sub(r"[年月日]$", "", str(value))
            consumed.add(index)
    if not parts:
        return None, set()
    ordered = [parts.get("year"), parts.get("month"), parts.get("day")]
    if parts.get("year"):
        formatted = f"{parts['year']}年"
    else:
        formatted = ""
    if parts.get("month"):
        formatted += f"{parts['month']}月"
    if parts.get("day"):
        formatted += f"{parts['day']}日"
    if all(parts.get(part) for part in ("year", "month", "day")):
        try:
            return date(int(parts["year"]), int(parts["month"]), int(parts["day"])), consumed
        except (TypeError, ValueError):
            pass
    return formatted or "/".join(value for value in ordered if value), consumed


def _combined_identity(answers: list[dict[str, Any]]) -> tuple[str | None, str | None, set[int]]:
    participant_id: str | None = None
    participant_name: str | None = None
    consumed: set[int] = set()
    for index, answer in enumerate(answers):
        explicit_id = answer.get("participant_id")
        explicit_name = answer.get("participant_name")
        if explicit_id not in (None, "") and participant_id is None:
            participant_id = str(explicit_id)
        if explicit_name not in (None, "") and participant_name is None:
            participant_name = str(explicit_name)
        normalized_key = _normalise(answer.get("normalized_key") or answer.get("question_id"))
        value = _cell_text(answer.get("value"))
        if value in (None, ""):
            continue
        label = _normalise(answer.get("label_japanese"))
        if participant_id is None and normalized_key in {_normalise("participant.id"), _normalise("ID")}:
            participant_id = str(value)
        if participant_name is None and normalized_key in {_normalise("participant.name"), _normalise("name")}:
            participant_name = str(value)
        if participant_id is None and label == _normalise("ID"):
            participant_id = str(value)
        if "氏名id" in label or normalized_key == _normalise("q4_admin_name"):
            text = str(value)
            match = re.search(r"ID\s*[:：]?\s*([A-Za-z0-9-]{4,})", text, flags=re.IGNORECASE)
            if match:
                if participant_id is None:
                    participant_id = match.group(1)
                before_id = re.split(r"\bID\b", text, maxsplit=1, flags=re.IGNORECASE)[0]
                before_id = re.sub(r"[/|]+", " ", before_id)
                before_id = re.sub(r"\s+", " ", before_id).strip(" '\"「『:：,，。")
                redaction_placeholder = re.search(
                    r"黒塗|墨消|匿名|redact|mask|[█■＊*●]{2,}", before_id, flags=re.IGNORECASE,
                )
                if participant_name is None and before_id and not redaction_placeholder:
                    participant_name = before_id
                consumed.add(index)
        if participant_id is None and "id" in label:
            match = re.search(r"\bID\s*[:：]?\s*([A-Za-z0-9-]{4,})", str(value), flags=re.IGNORECASE)
            if match:
                participant_id = match.group(1)
        if participant_name is None and (label == _normalise("氏名") or label == _normalise("名前")):
            participant_name = str(value)
    return participant_id, participant_name, consumed


def _append_unique(values: list[Any], value: Any) -> None:
    if value in (None, ""):
        return
    normalized = str(value)
    if all(str(existing) != normalized for existing in values):
        values.append(value)


def _load_records(records_dir: Path) -> list[tuple[Path, dict[str, Any]]]:
    paths = sorted(Path(records_dir).glob("*.json"), key=lambda item: _natural_key(item.stem))
    records: list[tuple[Path, dict[str, Any]]] = []
    for path in paths:
        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            raise ValueError(f"解析結果JSON必须是对象: {path.name}")
        records.append((path, payload))
    return records


def _style_data_row(sheet: Any, row: int, template_row: int = DATA_START_ROW) -> None:
    if row == template_row:
        return
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(template_row, column)
        destination = sheet.cell(row, column)
        if source.has_style:
            destination._style = copy(source._style)
        if source.number_format:
            destination.number_format = source.number_format
        if source.alignment:
            destination.alignment = copy(source.alignment)
    if sheet.row_dimensions[template_row].height is not None:
        sheet.row_dimensions[row].height = sheet.row_dimensions[template_row].height


def _write_cell(sheet: Any, row: int, column: int, value: Any) -> None:
    cell = sheet.cell(row, column)
    cell.value = _cell_text(value)
    if isinstance(cell.value, str):
        # Extracted strings are always data, even if they begin with '='.
        cell.data_type = "s"


def _append_unmapped_sheet(workbook: Any, rows: list[list[Any]]) -> None:
    if not rows:
        return
    if UNMAPPED_SHEET_NAME in workbook.sheetnames:
        del workbook[UNMAPPED_SHEET_NAME]
    sheet = workbook.create_sheet(UNMAPPED_SHEET_NAME)
    headers = ["ID", "JSON項目順", "ページ", "フィールドコード", "質問ID", "小項目ID", "項目", "値", "状態", "要確認"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="355C7D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append([_cell_text(value) for value in row])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {1: 18, 2: 14, 3: 10, 4: 24, 5: 34, 6: 20, 7: 48, 8: 48, 9: 18, 10: 12}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def write_excel_export(
    records_dir: Path,
    output_path: Path = DEFAULT_EXPORT_PATH,
    template_path: Path = TEMPLATE_PATH,
) -> dict[str, Any]:
    """Build the template workbook, retaining every non-context answer in the workbook."""
    records_dir = Path(records_dir)
    output_path = Path(output_path)
    template_path = Path(template_path)
    if not template_path.is_file():
        raise FileNotFoundError(f"指定格式Excel模板不存在: {template_path}")
    records = _load_records(records_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with _EXPORT_LOCK:
        workbook = load_workbook(template_path)
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Excel模板缺少工作表: {SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        field_keys = [sheet.cell(HEADER_ROW, column).value for column in range(1, sheet.max_column + 1)]
        field_labels = [sheet.cell(LABEL_ROW, column).value for column in range(1, sheet.max_column + 1)]
        if not all(isinstance(key, str) and key.strip() for key in field_keys):
            raise ValueError("Excel模板第2行必须为每一列提供字段代码。")
        if len({_normalise(key) for key in field_keys}) != len(field_keys):
            raise ValueError("Excel模板第2行存在重复字段代码，无法安全映射。")
        fields = [(str(key), str(label or "")) for key, label in zip(field_keys, field_labels)]
        column_by_key = {str(key): index + 1 for index, key in enumerate(field_keys)}
        measurement_counts: defaultdict[str, int] = defaultdict(int)
        unmapped_rows: list[list[Any]] = []
        unmatched_count = 0
        review_count_total = 0
        for output_row, (source_path, record) in enumerate(records, start=DATA_START_ROW):
            _style_data_row(sheet, output_row)
            raw_answers = record.get("answers")
            answers = [item for item in raw_answers if isinstance(item, dict)] if isinstance(raw_answers, list) else []
            record_id = _record_id(record, source_path)
            extracted_id, extracted_name, identity_consumed = _combined_identity(answers)
            participant_id = str(record.get("participant_id") or extracted_id or record_id)
            measurement_counts[participant_id] += 1
            row_values: dict[str, list[Any]] = {key: [] for key, _ in fields}
            matched_indexes: set[int] = set(identity_consumed)

            birth_value, birth_indexes = _combine_component_date(answers, birth=True)
            measured_date, date_indexes = _combine_component_date(answers, birth=False)
            if birth_value:
                row_values.setdefault("date of birth", []).append(birth_value)
                matched_indexes.update(birth_indexes)
            if measured_date:
                row_values.setdefault("date", []).append(measured_date)
                matched_indexes.update(date_indexes)

            for answer_index, answer in enumerate(answers):
                if answer_index in matched_indexes or _is_context_answer(answer):
                    continue
                split_regions = _selected_region_values(answer, fields)
                if split_regions is not None:
                    for key, value in split_regions.items():
                        row_values.setdefault(key, []).append(value)
                    matched_indexes.add(answer_index)
                    continue
                field_key = _find_field(answer, fields)
                if field_key is not None:
                    matched_indexes.add(answer_index)
                    _append_unique(row_values.setdefault(field_key, []), _answer_value(answer))

            # Identity values are kept in their dedicated columns. Combined
            # first-page identity strings are split only when an explicit ID is present.
            row_values.setdefault("ID", []).append(participant_id)
            if extracted_name:
                row_values.setdefault("name", []).append(extracted_name)
            participant_name = str(record.get("participant_name") or extracted_name or "")
            if participant_name:
                row_values.setdefault("name", []).append(participant_name)
            if record.get("label") not in (None, ""):
                row_values.setdefault("名前", []).append(record["label"])
            row_values.setdefault("Number of measurements", []).append(measurement_counts[participant_id])
            if record.get("measurement_date") and not row_values.get("date"):
                row_values.setdefault("date", []).append(_excel_date(record["measurement_date"]))

            review_count = sum(
                1 for answer in answers
                if answer.get("needs_review") or _answer_status(answer) in {"unclear", "conflict"}
            )
            review_count_total += review_count
            unmatched = [
                (index, answer) for index, answer in enumerate(answers)
                if index not in matched_indexes and not _is_context_answer(answer)
            ]
            unmatched_count += len(unmatched)
            for index, answer in unmatched:
                unmapped_rows.append([
                    participant_id,
                    index + 1,
                    answer.get("logical_page"),
                    answer.get("normalized_key"),
                    answer.get("question_id"),
                    answer.get("subquestion_id"),
                    answer.get("label_japanese") or "",
                    answer.get("value"),
                    _answer_status(answer),
                    bool(answer.get("needs_review") or _answer_status(answer) in {"unclear", "conflict"}),
                ])

            notes = []
            if review_count:
                notes.append(f"要確認{review_count}項目")
            if unmatched:
                notes.append(f"未対応{len(unmatched)}項目は「{UNMAPPED_SHEET_NAME}」シート")
            row_values.setdefault("備考", []).extend(notes)

            for field_key, values in row_values.items():
                column = column_by_key.get(field_key)
                if column is None:
                    continue
                nonempty = [value for value in values if value not in (None, "")]
                if not nonempty:
                    continue
                unique_values: list[Any] = []
                for value in nonempty:
                    _append_unique(unique_values, value)
                cell_value: Any = unique_values[0] if len(unique_values) == 1 else unique_values
                _write_cell(sheet, output_row, column, cell_value)
                if field_key in {"date", "date of birth"} and isinstance(sheet.cell(output_row, column).value, date):
                    sheet.cell(output_row, column).number_format = "yyyy/mm/dd"

        _append_unmapped_sheet(workbook, unmapped_rows)
        temporary_path: str | None = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=output_path.parent, suffix=".xlsx", delete=False,
            ) as temporary_file:
                temporary_path = temporary_file.name
            workbook.save(temporary_path)
            os.replace(temporary_path, output_path)
        finally:
            if temporary_path and Path(temporary_path).exists():
                Path(temporary_path).unlink()

    return {
        "excel_export_path": str(output_path),
        "excel_export_url": "/api/export/excel",
        "excel_record_count": len(records),
        "excel_unmapped_answer_count": unmatched_count,
        "excel_review_answer_count": review_count_total,
    }
